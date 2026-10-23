import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = ["title", "price", "currency", "rating", "in_stock", "url", "image_url"]
RU_HEADERS = ["Название", "Цена", "Валюта", "Рейтинг", "В наличии", "Ссылка", "Картинка"]


def to_csv(products, path: str) -> str:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig + ';' — файл корректно открывается двойным кликом в русском Excel
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(RU_HEADERS)
        for p in products:
            d = p.to_dict()
            writer.writerow([d[h] for h in HEADERS])
    return path


def to_xlsx(products, path: str, changes=None) -> str:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(RU_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for p in products:
        d = p.to_dict()
        ws.append([d[h] for h in HEADERS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate([48, 10, 8, 9, 11, 60, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    if changes:
        cs = wb.create_sheet("Изменения цен")
        cs.append(["Название", "Было", "Стало", "Разница", "Ссылка"])
        for c in changes:
            cs.append([c["title"], c["old_price"], c["new_price"], c["delta"], c["url"]])
    wb.save(path)
    return path
