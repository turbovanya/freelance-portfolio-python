from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..models import Product, RunResult
from .common import COLUMNS, csv_safe, display_value

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
UP_FILL = PatternFill("solid", fgColor="FDE2E1")
DOWN_FILL = PatternFill("solid", fgColor="E2F5E6")
PRICE_FORMAT = "#,##0.00"


def _sheet(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _link(ws: Worksheet, row: int, col: int) -> None:
    cell = ws.cell(row=row, column=col)
    if isinstance(cell.value, str) and cell.value.startswith("http"):
        cell.hyperlink, cell.style = cell.value, "Hyperlink"


def export_xlsx(result: RunResult, products: list[Product], path: Path) -> Path:
    path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    summary = wb.active
    summary.title = "Сводка"
    summary.append(["Отчёт о ценах", result.profile])
    summary["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Дата запуска", result.started_at.strftime("%d.%m.%Y %H:%M")),
        ("Страниц обработано", result.pages),
        ("Товаров собрано", len(result.products)),
        ("Товаров в выгрузке (после фильтров)", len(products)),
        ("Подорожало", len(result.increased)),
        ("Подешевело", len(result.decreased)),
        ("Новых товаров", len(result.new_products)),
        ("Пропало из каталога", len(result.removed)),
    ]
    summary.append([])
    for row in rows:
        summary.append(list(row))
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 24

    ws = wb.create_sheet("Товары")
    _sheet(ws, [title for _, title in COLUMNS], [52, 11, 9, 10, 11, 60, 60])
    for p in products:
        ws.append([csv_safe(display_value(p, key)) for key, _ in COLUMNS])
        ws.cell(row=ws.max_row, column=2).number_format = PRICE_FORMAT
        _link(ws, ws.max_row, 6)
    ws.auto_filter.ref = ws.dimensions

    if result.changes:
        cs = wb.create_sheet("Изменения цен")
        _sheet(
            cs, ["Название", "Было", "Стало", "Разница", "Изменение, %", "Ссылка"], [52, 11, 11, 11, 14, 60]
        )
        for c in result.changes:
            cs.append([csv_safe(c.title), c.old_price, c.new_price, c.delta, c.percent / 100, c.url])
            row = cs.max_row
            for col in range(1, 7):
                cs.cell(row=row, column=col).fill = UP_FILL if c.is_increase else DOWN_FILL
            for col in (2, 3, 4):
                cs.cell(row=row, column=col).number_format = PRICE_FORMAT
            cs.cell(row=row, column=5).number_format = "+0.0%;-0.0%"
            _link(cs, row, 6)
        cs.auto_filter.ref = cs.dimensions

    if result.new_products:
        ns = wb.create_sheet("Новые товары")
        _sheet(ns, ["Название", "Цена", "Ссылка"], [52, 11, 60])
        for p in result.new_products:
            ns.append([csv_safe(p.title), p.price, p.url])
            _link(ns, ns.max_row, 3)

    wb.save(path)
    return path
