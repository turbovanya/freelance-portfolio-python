import csv
from datetime import timedelta

import httpx
from openpyxl import load_workbook
from typer.testing import CliRunner

from conftest import NOW, product
from price_parser import cli
from price_parser.models import RunResult
from price_parser.notify import build_message, send_telegram
from price_parser.pipeline import RunOptions, run
from price_parser.storage import PriceStore


def test_full_run_over_two_pages(settings, profile, site_client):
    with PriceStore(settings.db_path) as store:
        result, selected, files = run(settings, profile, RunOptions(max_pages=0), site_client, store, now=NOW)

    assert result.pages == 2
    # Tipping the Velvet есть на обеих страницах, но в итоге один раз
    assert sorted(p.title for p in result.products) == [
        "A Light in the Attic",
        "Soumission",
        "Tipping the Velvet",
    ]
    assert result.first_run and len(selected) == 3
    assert {f: p.exists() for f, p in files.items()} == {"xlsx": True, "csv": True, "html": True}


def test_max_pages_limit(settings, profile, site_client):
    with PriceStore(settings.db_path) as store:
        result, _, _ = run(settings, profile, RunOptions(max_pages=1, formats=("csv",)), site_client, store)
    assert result.pages == 1 and len(result.products) == 2


def test_exports_content(settings, profile, site_client):
    with PriceStore(settings.db_path) as store:
        store.save_run(
            profile.name,
            NOW - timedelta(days=1),
            [
                product(
                    "https://books.toscrape.com/a-light-in-the-attic_1000/index.html",
                    45.00,
                    title="A Light in the Attic",
                )
            ],
        )
        result, _, files = run(settings, profile, RunOptions(max_pages=0), site_client, store, now=NOW)

    assert len(result.changes) == 1 and result.changes[0].is_increase

    wb = load_workbook(files["xlsx"])
    assert wb.sheetnames == ["Сводка", "Товары", "Изменения цен", "Новые товары"]
    assert wb["Товары"].max_row == 4
    assert wb["Изменения цен"]["C2"].value == 51.77

    rows = list(csv.reader(files["csv"].open(encoding="utf-8-sig"), delimiter=";"))
    assert rows[0][:2] == ["Название", "Цена"] and len(rows) == 4

    html = files["html"].read_text(encoding="utf-8")
    assert "A Light in the Attic" in html and "+6.77" in html


def test_csv_escapes_formulas(settings, store):
    from price_parser.exporters import export_csv

    result = RunResult("t", NOW, [])
    path = export_csv(result, [product(title='=HYPERLINK("http://x")')], settings.output_dir / "p")
    row = list(csv.reader(path.open(encoding="utf-8-sig"), delimiter=";"))[1]
    assert row[0].startswith("'=")


def test_telegram_message_and_send(settings):
    from price_parser.models import PriceChange

    result = RunResult("Books", NOW, [], changes=[PriceChange("<Книга>", "https://x/1", 10, 12)])
    text = build_message(result)
    assert "&lt;Книга&gt;" in text and "+20.0%" in text
    assert build_message(result, threshold_percent=50) is None

    sent = []
    transport = httpx.MockTransport(lambda r: sent.append(r) or httpx.Response(200, json={"ok": True}))
    settings = settings.model_copy(update={"telegram_bot_token": "1:x", "telegram_chat_id": "42"})
    assert send_telegram(settings, text, transport=transport)
    assert sent[0].url.path == "/bot1:x/sendMessage"


def test_cli_run(monkeypatch, settings, site_client):
    monkeypatch.setattr(cli, "Settings", lambda: settings)
    monkeypatch.setattr(cli, "HttpClient", lambda _s: site_client)
    result = CliRunner().invoke(cli.app, ["run", "--max-pages", "0", "--format", "csv", "--min-rating", "3"])
    assert result.exit_code == 0, result.output
    assert "страница 2" in result.output
    assert (settings.output_dir / "products.csv").exists()


def test_cli_rejects_unknown_format(monkeypatch, settings):
    monkeypatch.setattr(cli, "Settings", lambda: settings)
    result = CliRunner().invoke(cli.app, ["run", "--format", "pdf"])
    assert result.exit_code != 0
