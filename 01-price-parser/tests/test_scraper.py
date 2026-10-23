import pytest
import requests
from openpyxl import load_workbook

from parser.export import to_csv, to_xlsx
from parser.config import Settings
from parser.scraper import Product, fetch, filter_products, parse_catalog_page, parse_price
from parser.storage import PriceStore

HTML = """
<html><body><ol>
<li><article class="product_pod">
  <div class="image_container"><a href="a-light-in-the-attic_1000/index.html"><img src="../media/a.jpg"></a></div>
  <p class="star-rating Three"></p>
  <h3><a href="a-light-in-the-attic_1000/index.html" title="A Light in the Attic">A Light in the ...</a></h3>
  <div class="product_price"><p class="price_color">Â£51.77</p>
  <p class="instock availability"><i class="icon-ok"></i> In stock</p></div>
</article></li>
<li><article class="product_pod">
  <p class="star-rating Five"></p>
  <h3><a href="tipping_999/index.html" title="Tipping the Velvet">Tipping</a></h3>
  <p class="price_color">£53.74</p>
  <p class="availability">Out of stock</p>
</article></li>
<li><article class="product_pod"><h3>broken card without price</h3></article></li>
</ol><ul class="pager"><li class="next"><a href="page-2.html">next</a></li></ul></body></html>
"""
URL = "https://books.toscrape.com/catalogue/page-1.html"


def sample(price=10.0, rating=3, stock=True, url="u1"):
    return Product("T", price, "£", rating, stock, url, "")


def test_parse_price_variants():
    assert parse_price("£51.77") == (51.77, "£")
    assert parse_price("Â£1 200,50") == (1200.5, "£")
    with pytest.raises(ValueError):
        parse_price("free")


def test_parse_catalog_page():
    products, nxt = parse_catalog_page(HTML, URL)
    assert len(products) == 2  # битая карточка пропущена
    first, second = products
    assert first.title == "A Light in the Attic"
    assert first.price == 51.77 and first.currency == "£"
    assert first.rating == 3 and first.in_stock is True
    assert first.url == "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html"
    assert second.in_stock is False and second.rating == 5
    assert nxt == "https://books.toscrape.com/catalogue/page-2.html"


def test_last_page_has_no_next():
    _, nxt = parse_catalog_page("<html></html>", URL)
    assert nxt is None


def test_filters():
    items = [sample(5, 1, True, "a"), sample(20, 4, False, "b"), sample(50, 5, True, "c")]
    assert [p.url for p in filter_products(items, min_price=10)] == ["b", "c"]
    assert [p.url for p in filter_products(items, max_price=20, min_rating=4)] == ["b"]
    assert [p.url for p in filter_products(items, only_in_stock=True)] == ["a", "c"]


def test_price_change_detection():
    store = PriceStore(":memory:")
    assert store.save_and_diff([sample(10.0)]) == []  # первый запуск: не с чем сравнивать
    assert store.save_and_diff([sample(10.0)]) == []  # цена не менялась
    changes = store.save_and_diff([sample(12.5)])
    assert len(changes) == 1 and changes[0]["delta"] == 2.5
    assert store.last_prices()["u1"] == 12.5


def test_exports(tmp_path):
    items = [sample(10.0), sample(20.0, url="u2")]
    csv_path = to_csv(items, str(tmp_path / "p.csv"))
    lines = open(csv_path, encoding="utf-8-sig").read().splitlines()
    assert lines[0].startswith("Название;Цена")
    assert len(lines) == 3
    changes = [{"title": "T", "old_price": 1, "new_price": 2, "delta": 1, "url": "u"}]
    wb = load_workbook(to_xlsx(items, str(tmp_path / "p.xlsx"), changes))
    assert wb.sheetnames == ["Товары", "Изменения цен"]
    assert wb["Товары"].max_row == 3
    assert wb["Товары"]["E2"].value == "да"


class FakeResponse:
    def __init__(self, status, text="ok", content_type="text/html"):
        self.status_code, self.text, self.encoding = status, text, None
        self.headers = {"Content-Type": content_type}

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(response=self)


class FakeSession:
    def __init__(self, *responses):
        self.responses, self.calls = list(responses), 0

    def get(self, url, timeout):
        self.calls += 1
        r = self.responses.pop(0)
        if isinstance(r, Exception):
            raise r
        return r


FAST = Settings(retries=3, request_delay=0)


def test_fetch_retries_on_server_and_network_errors():
    session = FakeSession(requests.ConnectionError("down"), FakeResponse(503), FakeResponse(200, "page"))
    assert fetch(session, URL, FAST) == "page"
    assert session.calls == 3


def test_fetch_does_not_retry_404():
    session = FakeSession(FakeResponse(404), FakeResponse(200))
    with pytest.raises(requests.HTTPError):
        fetch(session, URL, FAST)
    assert session.calls == 1


def test_fetch_gives_up_after_retries():
    session = FakeSession(*[FakeResponse(500)] * 3)
    with pytest.raises(RuntimeError):
        fetch(session, URL, FAST)
